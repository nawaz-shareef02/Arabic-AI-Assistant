import os
import docx
from app.services.parsers.base import BaseParser, ParsingResult, DocumentParsingError

class DOCXParser(BaseParser):
    def parse(self, file_path: str) -> ParsingResult:
        if not os.path.exists(file_path):
            raise DocumentParsingError("File not found")
            
        if os.path.getsize(file_path) == 0:
            raise DocumentParsingError("Empty document")
            
        try:
            doc = docx.Document(file_path)
            text_parts = []
            
            # Extract paragraphs
            for para in doc.paragraphs:
                if para.text:
                    text_parts.append(para.text)
                    
            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text:
                            text_parts.append(cell.text)
                            
            text = "\n".join(text_parts)
        except Exception as e:
            raise DocumentParsingError("Invalid DOCX: Failed to read XML document structure", str(e))
            
        if not text.strip():
            raise DocumentParsingError("Empty document")
            
        return ParsingResult(text=text, page_count=None, char_count=len(text))
