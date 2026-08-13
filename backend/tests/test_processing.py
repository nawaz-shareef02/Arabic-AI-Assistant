import pytest
import os
import tempfile
from unittest.mock import MagicMock, patch
import docx

from app.models.document import Document
from app.models.parsed_document import ParsedDocument
from app.services.parsers.base import DocumentParsingError
from app.services.parsers.txt import TXTParser
from app.services.parsers.markdown import MarkdownParser
from app.services.parsers.pdf import PDFParser
from app.services.parsers.docx import DOCXParser
from app.services.parsers.factory import ParserFactory
from app.services.parser_service import ParserService, normalize_text, detect_language_and_confidence
from app.services.document_service import DocumentService
from app.schemas.document import DocumentStatus

# 1. Base DB Setup Fixtures
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from app.database.base import Base

SQLALCHEMY_DATABASE_URL = "sqlite://"

engine = create_engine(
    SQLALCHEMY_DATABASE_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)

@pytest.fixture(name="db")
def fixture_db():
    Base.metadata.create_all(bind=engine)
    session = TestingSessionLocal()
    try:
        yield session
    finally:
        session.close()
        Base.metadata.drop_all(bind=engine)

# 2. Text Normalization tests
def test_normalization():
    raw_text = "Hello \x00\x07 world!   This   has   too   many   spaces. \n\n\n\nNew paragraph."
    normalized = normalize_text(raw_text)
    assert "Hello  world!" not in normalized  # control characters removed
    assert "This has too many spaces." in normalized  # spacing clean
    assert "\n\n\n" not in normalized  # consecutive blank lines limited to 2

# 3. Language detection & confidence tests
def test_language_detection():
    english_text = "This is a purely English text for checking the language ratio."
    lang, conf = detect_language_and_confidence(english_text)
    assert lang == "English"
    assert conf > 0.8

    arabic_text = "هذا نص باللغة العربية الفصحى لاختبار مدى دقة الكشف."
    lang, conf = detect_language_and_confidence(arabic_text)
    assert lang == "Arabic"
    assert conf > 0.8

    mixed_text = "العربية Arabic text language"
    lang, conf = detect_language_and_confidence(mixed_text)
    assert lang == "Mixed"
    assert conf < 1.0

# 4. TXT parser tests
def test_txt_parser():
    parser = TXTParser()
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as tmp:
        tmp.write("Hello plain text".encode("utf-8"))
        tmp_path = tmp.name

    try:
        res = parser.parse(tmp_path)
        assert res.text == "Hello plain text"
        assert res.char_count == len("Hello plain text")
        assert res.page_count is None
    finally:
        os.remove(tmp_path)

def test_txt_parser_empty():
    parser = TXTParser()
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        with pytest.raises(DocumentParsingError) as exc_info:
            parser.parse(tmp_path)
        assert "Empty document" in str(exc_info.value)
    finally:
        os.remove(tmp_path)

# 5. Markdown parser tests
def test_markdown_parser():
    parser = MarkdownParser()
    with tempfile.NamedTemporaryFile(suffix=".md", delete=False) as tmp:
        tmp.write("# Title\n- Bullet point".encode("utf-8"))
        tmp_path = tmp.name

    try:
        res = parser.parse(tmp_path)
        assert "# Title" in res.text
    finally:
        os.remove(tmp_path)

# 6. PDF parser tests with mocking
def test_pdf_parser_success():
    parser = PDFParser()
    mock_page = MagicMock()
    mock_page.extract_text.return_value = "PDF page text content"
    mock_reader = MagicMock()
    mock_reader.pages = [mock_page]
    mock_reader.is_encrypted = False

    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
        tmp.write(b"%PDF-1.4 ...")
        tmp_path = tmp.name

    try:
        with patch("app.services.parsers.pdf.PdfReader", return_value=mock_reader):
            res = parser.parse(tmp_path)
            assert res.text == "PDF page text content"
            assert res.page_count == 1
    finally:
        os.remove(tmp_path)

def test_pdf_parser_corrupted():
    parser = PDFParser()
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
        tmp.write(b"invalid pdf structure")
        tmp_path = tmp.name

    try:
        with pytest.raises(DocumentParsingError) as exc_info:
            parser.parse(tmp_path)
        assert "Corrupted PDF" in str(exc_info.value)
    finally:
        os.remove(tmp_path)

# 7. DOCX parser tests
def test_docx_parser_success():
    parser = DOCXParser()
    doc = docx.Document()
    doc.add_paragraph("DOCX paragraph text.")
    
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = tmp.name
    
    try:
        doc.save(tmp_path)
        res = parser.parse(tmp_path)
        assert res.text == "DOCX paragraph text."
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)

def test_docx_parser_corrupted():
    parser = DOCXParser()
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp.write(b"corrupted zip archive")
        tmp_path = tmp.name

    try:
        with pytest.raises(DocumentParsingError) as exc_info:
            parser.parse(tmp_path)
        assert "Invalid DOCX" in str(exc_info.value)
    finally:
        os.remove(tmp_path)

# 8. Parser Factory test
def test_parser_factory():
    assert isinstance(ParserFactory.get_parser("file.pdf"), PDFParser)
    assert isinstance(ParserFactory.get_parser("file.docx"), DOCXParser)
    assert isinstance(ParserFactory.get_parser("file.txt"), TXTParser)
    assert isinstance(ParserFactory.get_parser("file.md"), MarkdownParser)
    
    with pytest.raises(DocumentParsingError):
        ParserFactory.get_parser("file.exe")

# 9. Lifecycle processing test (DocumentService integration)
def test_document_processing_lifecycle(db):
    from app.models.user import User
    from app.models.knowledge_base import KnowledgeBase
    import uuid as py_uuid
    
    # Setup mock user and KB
    user = User(email="test@user.com", hashed_password="hashed_password", full_name="Tester", role="employee", organization="Test Org")
    db.add(user)
    db.commit()
    
    kb = KnowledgeBase(name="Test KB", owner_id=user.id)
    db.add(kb)
    db.commit()
    
    # Create temp text file
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as tmp:
        tmp.write("This is a simple document containing English words.".encode("utf-8"))
        tmp_path = tmp.name
        
    doc = Document(
        uuid=py_uuid.uuid4(),
        knowledge_base_id=kb.id,
        filename="test.txt",
        storage_path=tmp_path,
        mime_type="text/plain",
        file_size=os.path.getsize(tmp_path),
        status="Uploaded",
        created_by=user.id,
        updated_by=user.id
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    
    from app.tasks.indexing_tasks import process_document_async
    
    # Run pipeline process synchronously using a mock SessionLocal to point to our in-memory SQLite DB
    with patch("app.tasks.indexing_tasks.SessionLocal", return_value=db):
        with patch.object(db, "close", return_value=None):
            with patch("app.tasks.indexing_tasks.IndexingService") as mock_indexing:
                mock_indexing.return_value.index_document.return_value = 1
                with patch("app.tasks.indexing_tasks.run_intelligence_pipeline") as mock_intel:
                    mock_intel.delay = MagicMock()
                    process_document_async.run(document_id=doc.id, knowledge_base_id=kb.id)
        
    db.refresh(doc)
    assert doc.status == "Parsed"
    assert doc.language == "English"
    assert doc.parsed_document is not None
    assert doc.parsed_document.parsed_text == "This is a simple document containing English words."
    assert doc.parsed_document.char_count == len("This is a simple document containing English words.")
    assert doc.parsed_document.language_confidence > 0.8
    assert doc.parsed_document.processing_duration >= 0
    
    os.remove(tmp_path)
